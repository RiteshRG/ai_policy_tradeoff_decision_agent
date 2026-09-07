from langchain_core.tools import tool
from pptx import Presentation
from pptx.util import Inches, Pt
import json
import os
import tempfile
import shutil

DEFAULT_WEIGHTS = {
    "food_water": 0.35,
    "energy": 0.35,
    "vehicle_consumer": 0.30
}

@tool
def score_policy_options(findings_json: str, weights_json: str = "") -> str:
    """Calculate deterministic weighted scores for policy options.
    findings_json: JSON string of {option: {criterion: value}}.
    weights_json: optional JSON string of {criterion: weight}; if
    omitted or empty, default weights are used."""
    findings = json.loads(findings_json)
    weights = json.loads(weights_json) if weights_json else DEFAULT_WEIGHTS

    results = {}
    for option, scores in findings.items():
        total = sum(scores[c] * weights[c] for c in weights)
        results[option] = round(total, 2)
    return json.dumps(results)


@tool
def generate_decision_ppt(report_json: str) -> str:
    """Create exactly one slide summarizing the decision report.
    report_json: a JSON STRING (not a dict object) with keys:
    title (str), recommendation (str), scores (dict of option->score),
    trade_offs (list of short strings), conflicts (list of short
    strings), caveats (list of short strings), decision_conditions
    (list of short strings), sources (list of short strings)."""
    data = json.loads(report_json)

    title_text = data.get("title", "India E20 Policy Trade-off Analysis")
    recommendation = data.get("recommendation")
    if not recommendation:
        raise ValueError("Missing 'recommendation' in report_json")

    scores = data.get("scores", {})
    trade_offs = data.get("trade_offs", [])
    conflicts = data.get("conflicts", [])
    caveats = data.get("caveats", [])
    decision_conditions = data.get("decision_conditions", [])
    sources = data.get("sources", [])

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout

    def add_text(left, top, width, height, text, size=14, bold=False):
        box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        tf = box.text_frame
        tf.word_wrap = True
        tf.text = text
        for p in tf.paragraphs:
            p.font.size = Pt(size)
            p.font.bold = bold
        return box

    # Title
    add_text(0.5, 0.25, 12.3, 0.7, title_text, size=26, bold=True)

    # Recommendation (prominent)
    add_text(0.5, 1.05, 12.3, 0.6, f"Recommended: {recommendation}", size=18, bold=True)

    # Scores (compact line)
    if scores:
        score_line = " | ".join(f"{opt}: {val}" for opt, val in scores.items())
        add_text(0.5, 1.75, 12.3, 0.5, f"Deterministic Scores — {score_line}", size=13)

    # Trade-offs
    if trade_offs:
        add_text(0.5, 2.4, 6.0, 1.6, "Key Trade-offs:\n" + "\n".join(f"- {t}" for t in trade_offs), size=12)

    # Conflicts / caveats
    combined_notes = []
    if conflicts:
        combined_notes.append("Conflicts:\n" + "\n".join(f"- {c}" for c in conflicts))
    if caveats:
        combined_notes.append("Caveats:\n" + "\n".join(f"- {c}" for c in caveats))
    if combined_notes:
        add_text(6.7, 2.4, 6.1, 1.6, "\n\n".join(combined_notes), size=12)

    # Decision conditions
    if decision_conditions:
        add_text(0.5, 4.2, 12.3, 1.2, "Applies under: " + "; ".join(decision_conditions), size=12)

    # Sources (small, bottom)
    if sources:
        add_text(0.5, 6.7, 12.3, 0.6, "Sources: " + "; ".join(sources), size=9)



    os.makedirs("output", exist_ok=True)
    output_path = "output/final_decision.pptx"

    # Save to a temp file first, then move into place — guarantees the
    # final path never contains a half-written/corrupted file.
    tmp_fd, tmp_path = tempfile.mkstemp(suffix=".pptx", dir="output")
    os.close(tmp_fd)
    prs.save(tmp_path)
    shutil.move(tmp_path, output_path)

    return output_path